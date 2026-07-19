import os

from datetime import datetime

from docx import Document
from docx.shared import Pt


OUTPUT_DIR = "generated_docs"

os.makedirs(OUTPUT_DIR, exist_ok=True)


def create_docx(content: str):

    document = Document()

    title = document.add_heading(
        "AI Generated Business Document",
        level=1
    )

    title.style.font.size = Pt(20)

    document.add_paragraph()

    document.add_heading("Generated On", level=2)

    document.add_paragraph(
        datetime.now().strftime("%d %B %Y %H:%M")
    )

    document.add_heading("Document", level=2)

    paragraphs = content.split("\n")

    for line in paragraphs:

        line = line.strip()

        if not line:
            continue

        # Markdown Heading
        if line.startswith("**") and line.endswith("**"):

            heading = line.replace("**", "").strip()

            document.add_heading(heading, level=2)

        elif line.startswith("# "):

            document.add_heading(
                line.replace("# ", "").strip(),
                level=1
            )

        elif line.startswith("## "):

            document.add_heading(
                line.replace("## ", "").strip(),
                level=2
            )

        elif line.startswith("### "):

            document.add_heading(
                line.replace("### ", "").strip(),
                level=3
            )

        elif line.startswith("- "):

            document.add_paragraph(
                line.replace("- ", "").strip(),
                style="List Bullet"
            )

        elif line.startswith("* "):

            document.add_paragraph(
                line.replace("* ", "").strip(),
                style="List Bullet"
            )

        elif line.startswith("•"):

            document.add_paragraph(
                line.replace("•", "").strip(),
                style="List Bullet"
            )

        else:

            document.add_paragraph(line)

    filename = f"document_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"

    filepath = os.path.join(
        OUTPUT_DIR,
        filename
    )

    document.save(filepath)

    return filepath